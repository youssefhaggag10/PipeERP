import os
from collections.abc import Generator
from datetime import UTC, datetime
from io import BytesIO
from secrets import token_urlsafe

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

TEST_PASSWORD = token_urlsafe(32)
VIEWER_PASSWORD = token_urlsafe(32)
os.environ.setdefault("SECRET_KEY", token_urlsafe(48))
os.environ.setdefault("APP_ENV", "test")

from app.infrastructure.database.base import Base  # noqa: E402
from app.infrastructure.database.session import get_database_session  # noqa: E402
from app.main import app  # noqa: E402
from app.modules.identity.models import User  # noqa: E402
from app.modules.identity.service import ClientContext, create_initial_admin  # noqa: E402
from app.modules.inventory.models import InventoryBalance, InventoryLayer  # noqa: E402
from app.modules.master_data.models import (  # noqa: E402
    CompanySettings,
    Partner,
    Product,
    UnitOfMeasure,
    Warehouse,
)
from app.modules.purchasing.models import PurchaseOrder, SupplierInvoice  # noqa: E402
from app.modules.sales.models import (  # noqa: E402
    CustomerInvoice,
    SalesOrder,
    SalesOrderLine,
    SalesQuotation,
    SalesQuotationLine,
)
from app.modules.treasury.models import (  # noqa: E402
    FinancialAccount,
    PaymentAllocation,
    PaymentTransaction,
)


def _database() -> sessionmaker[Session]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _client(factory: sessionmaker[Session]) -> TestClient:
    def override_database() -> Generator[Session, None, None]:
        with factory() as db:
            yield db

    app.dependency_overrides[get_database_session] = override_database
    return TestClient(app)


def _seed(factory: sessionmaker[Session]) -> dict[str, str]:
    now = datetime(2026, 8, 20, 10, 30, tzinfo=UTC)
    with factory.begin() as db:
        create_initial_admin(
            db,
            username="admin",
            display_name="مدير التقارير",
            password=TEST_PASSWORD,
            client=ClientContext("127.0.0.1", "test", "bootstrap"),
        )
        actor_id = db.scalar(select(User.id).where(User.normalized_username == "admin"))
        assert actor_id is not None
        unit = UnitOfMeasure(
            code="PCS",
            normalized_code="PCS",
            name_ar="قطعة",
            symbol="قطعة",
            decimal_places=0,
            is_active=True,
            version=1,
        )
        warehouse = Warehouse(
            code="MAIN",
            normalized_code="MAIN",
            name_ar="المصنع",
            is_default=True,
            is_active=True,
            version=1,
        )
        customer = Partner(
            code="CUS-RPT",
            normalized_code="CUS-RPT",
            name_ar="عميل التقارير",
            phone="01000000000",
            address="القاهرة",
            tax_number="",
            is_customer=True,
            is_supplier=False,
            is_active=True,
            version=1,
        )
        supplier = Partner(
            code="SUP-RPT",
            normalized_code="SUP-RPT",
            name_ar="مورد التقارير",
            phone="",
            address="",
            tax_number="",
            is_customer=False,
            is_supplier=True,
            is_active=True,
            version=1,
        )
        cash = FinancialAccount(
            code="CASH-RPT",
            normalized_code="CASH-RPT",
            name_ar="الخزينة الرئيسية",
            account_type="cash",
            opening_balance=0,
            is_default=True,
            is_active=True,
            notes="",
            version=1,
        )
        db.add_all([unit, warehouse, customer, supplier, cash])
        db.flush()
        product = Product(
            code="FG-RPT",
            normalized_code="FG-RPT",
            name_ar="ماسورة اختبار",
            product_type="finished_good",
            unit_id=unit.id,
            category_id=None,
            min_stock=0,
            track_lots=True,
            standard_weight_kg=0,
            weight_tolerance_percent=5,
            is_active=True,
            version=1,
        )
        db.add(product)
        db.flush()
        db.add(
            CompanySettings(
                id=1,
                company_name_ar="مصنع بايب",
                phone="0123456789",
                address="المنطقة الصناعية",
                tax_number="TAX-1",
                currency_code="EGP",
                currency_decimal_places=2,
                tax_enabled=False,
                default_tax_rate=0,
                default_warehouse_id=warehouse.id,
                version=1,
            )
        )
        order = SalesOrder(
            order_number="SO-RPT-1",
            customer_id=customer.id,
            warehouse_id=warehouse.id,
            billing_method="piece",
            status="delivered",
            order_date=now,
            notes="اختبار",
            subtotal=100,
            discount_amount=0,
            transport_amount=0,
            tax_amount=0,
            total=100,
            version=1,
            created_by_id=actor_id,
        )
        db.add(order)
        db.flush()
        db.add(
            SalesOrderLine(
                sales_order_id=order.id,
                product_id=product.id,
                quantity=2,
                unit="قطعة",
                unit_price=50,
                line_total=100,
                standard_weight_kg=0,
                billing_weight_kg=0,
                price_per_kg=0,
                notes="",
            )
        )
        invoice = CustomerInvoice(
            invoice_number="SI-RPT-1",
            sales_order_id=order.id,
            customer_id=customer.id,
            invoice_type="standard",
            status="posted",
            invoice_date=now,
            subtotal=100,
            discount_amount=0,
            transport_amount=0,
            tax_amount=0,
            total=100,
            notes="",
            posted_at=now,
            reversal_reason="",
            version=1,
        )
        db.add(invoice)
        purchase = PurchaseOrder(
            order_number="PO-RPT-1",
            supplier_id=supplier.id,
            warehouse_id=warehouse.id,
            status="received",
            order_date=now,
            notes="",
            total=80,
            version=1,
            created_by_id=actor_id,
        )
        db.add(purchase)
        db.flush()
        supplier_invoice = SupplierInvoice(
            invoice_number="PI-RPT-1",
            supplier_invoice_number="EXT-1",
            purchase_order_id=purchase.id,
            supplier_id=supplier.id,
            status="posted",
            total=80,
            posted_at=now,
            version=1,
        )
        db.add(supplier_invoice)
        quotation = SalesQuotation(
            quotation_number="QT-RPT-1",
            customer_id=customer.id,
            quotation_date=now,
            valid_until=now,
            status="sent",
            total=120,
            notes="عرض تجريبي",
            version=1,
            created_by_id=actor_id,
        )
        db.add(quotation)
        db.flush()
        db.add(
            SalesQuotationLine(
                quotation_id=quotation.id,
                product_id=product.id,
                item_name=product.name_ar,
                quantity=2,
                unit="قطعة",
                unit_price=60,
                line_total=120,
                notes="",
            )
        )
        payment = PaymentTransaction(
            transaction_number="RC-RPT-1",
            transaction_date=now,
            transaction_type="customer_receipt",
            partner_id=customer.id,
            financial_account_id=cash.id,
            amount=40,
            payment_method="cash",
            reference_type="sale",
            reference_id=order.id,
            customer_invoice_id=invoice.id,
            supplier_invoice_id=None,
            idempotency_key="reports-payment",
            request_hash="0" * 64,
            status="posted",
            notes="دفعة",
            posted_by_id=actor_id,
            reversal_reason="",
        )
        db.add(payment)
        db.flush()
        db.add(
            PaymentAllocation(
                payment_transaction_id=payment.id,
                customer_invoice_id=invoice.id,
                supplier_invoice_id=None,
                amount=40,
            )
        )
        db.add(
            InventoryBalance(
                product_id=product.id,
                warehouse_id=warehouse.id,
                quantity_on_hand=10,
                weight_on_hand_kg=0,
                version=1,
            )
        )
        db.add(
            InventoryLayer(
                product_id=product.id,
                warehouse_id=warehouse.id,
                lot_id=None,
                source_type="opening",
                source_id="RPT",
                source_line_id=None,
                cost_basis="quantity",
                quantity_received=10,
                quantity_remaining=10,
                weight_received_kg=0,
                weight_remaining_kg=0,
                unit_cost=8,
                received_at=now,
                version=1,
            )
        )
        db.flush()
        return {
            "invoice_id": str(invoice.id),
            "quotation_id": str(quotation.id),
            "customer_id": str(customer.id),
        }


def _login(client: TestClient) -> None:
    response = client.post(
        "/api/v1/auth/login", json={"username": "admin", "password": TEST_PASSWORD}
    )
    assert response.status_code == 200, response.text


def _csrf(client: TestClient) -> dict[str, str]:
    token = client.cookies.get("pipeerp_csrf")
    assert token is not None
    return {"X-CSRF-Token": token}


def test_reports_exports_and_a4_print_data() -> None:
    factory = _database()
    ids = _seed(factory)
    client = _client(factory)
    _login(client)
    period = "date_from=2026-01-01&date_to=2026-12-31"

    for key in (
        "sales",
        "purchases",
        "customer_balances",
        "supplier_balances",
        "payments",
        "inventory_valuation",
    ):
        response = client.get(f"/api/v1/reports/generate?report_key={key}&{period}")
        assert response.status_code == 200, f"{key}: {response.text}"
        assert response.json()["rows"]

    sales = client.get(f"/api/v1/reports/generate?report_key=sales&{period}").json()
    assert sales["summary"] == {
        "original": "100.00",
        "returned": "0.00",
        "net": "100.00",
        "paid": "40.00",
        "remaining": "60.00",
    }
    inventory = client.get(
        f"/api/v1/reports/generate?report_key=inventory_valuation&{period}"
    ).json()
    assert inventory["summary"]["inventory_value"] == "80.00"

    exported = client.get(f"/api/v1/reports/export.xlsx?report_key=sales&{period}")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(filename=BytesIO(exported.content))
    assert workbook.active.freeze_panes == "A5"
    assert workbook.active.auto_filter.ref is not None

    invoice = client.get(f"/api/v1/reports/print/sales-invoices/{ids['invoice_id']}")
    assert invoice.status_code == 200, invoice.text
    assert invoice.json()["document_title"] == "فاتورة مبيعات"
    assert invoice.json()["paid"] == "40.00"
    assert len(invoice.json()["lines"]) == 1

    quote = client.get(f"/api/v1/reports/print/quotations/{ids['quotation_id']}")
    assert quote.status_code == 200, quote.text
    assert quote.json()["document_title"] == "عرض سعر"

    with factory.begin() as db:
        actor_id = db.scalar(select(User.id).where(User.normalized_username == "admin"))
        customer_id = db.scalar(select(Partner.id).where(Partner.code == "CUS-RPT"))
        warehouse_id = db.scalar(select(Warehouse.id).where(Warehouse.code == "MAIN"))
        assert actor_id is not None and customer_id is not None and warehouse_id is not None
        db.add(
            SalesOrder(
                order_number="SO-DRAFT-RPT",
                customer_id=customer_id,
                warehouse_id=warehouse_id,
                billing_method="piece",
                status="draft",
                order_date=datetime(2026, 8, 20, 12, 0, tzinfo=UTC),
                notes="مسودة للمراجعة",
                subtotal=25,
                discount_amount=0,
                transport_amount=0,
                tax_amount=0,
                total=25,
                version=1,
                created_by_id=actor_id,
            )
        )

    statement = client.get(
        f"/api/v1/reports/print/customer-statements/{ids['customer_id']}?"
        f"{period}&detailed=true&include_drafts=true"
    )
    assert statement.status_code == 200, statement.text
    assert statement.json()["statement"]["closing_balance"] == "60.00"
    assert len(statement.json()["invoice_details"]["SI-RPT-1"]) == 1
    draft_line = next(
        line
        for line in statement.json()["statement"]["lines"]
        if line["document_number"] == "SO-DRAFT-RPT"
    )
    assert float(draft_line["debit"]) == 0
    assert "لا تدخل في الرصيد" in draft_line["notes"]
    statement_export = client.get(
        f"/api/v1/reports/export/customer-statements/{ids['customer_id']}.xlsx?"
        f"{period}&detailed=true&include_drafts=true"
    )
    assert statement_export.status_code == 200
    statement_book = load_workbook(filename=BytesIO(statement_export.content))
    assert statement_book.active.freeze_panes == "A6"

    invalid = client.get(
        "/api/v1/reports/generate?report_key=sales&date_from=2026-12-31&date_to=2026-01-01"
    )
    assert invalid.status_code == 422
    app.dependency_overrides.clear()


def test_report_permission_does_not_grant_source_document_access() -> None:
    factory = _database()
    ids = _seed(factory)
    with _client(factory) as admin:
        _login(admin)
        role = admin.post(
            "/api/v1/identity/roles",
            headers=_csrf(admin),
            json={
                "code": "report_viewer",
                "name_ar": "مشاهد التقارير",
                "permissions": ["reports.read"],
            },
        )
        assert role.status_code == 201, role.text
        viewer = admin.post(
            "/api/v1/identity/users",
            headers=_csrf(admin),
            json={
                "username": "report.viewer",
                "display_name": "مشاهد التقارير",
                "password": VIEWER_PASSWORD,
                "role_codes": ["report_viewer"],
                "must_change_password": False,
            },
        )
        assert viewer.status_code == 201, viewer.text

    with _client(factory) as viewer_client:
        login = viewer_client.post(
            "/api/v1/auth/login",
            json={"username": "report.viewer", "password": VIEWER_PASSWORD},
        )
        assert login.status_code == 200, login.text
        period = "date_from=2026-01-01&date_to=2026-12-31"
        assert (
            viewer_client.get(f"/api/v1/reports/generate?report_key=sales&{period}").status_code
            == 200
        )
        assert (
            viewer_client.get(
                f"/api/v1/reports/print/sales-invoices/{ids['invoice_id']}"
            ).status_code
            == 403
        )
        assert (
            viewer_client.get(
                f"/api/v1/reports/print/customer-statements/{ids['customer_id']}?{period}"
            ).status_code
            == 403
        )
    app.dependency_overrides.clear()
