from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Literal, cast
from uuid import UUID

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session

from app.domain.common.decimal import money, quantity
from app.modules.inventory.models import InventoryBalance, InventoryLayer
from app.modules.master_data.models import CompanySettings, Partner, Product, Warehouse
from app.modules.purchasing.models import PurchaseOrder, SupplierInvoice
from app.modules.reports.schemas import (
    CustomerStatementPrintView,
    CustomerStatementSummaryView,
    PrintCompanyView,
    PrintDocumentLineView,
    PrintDocumentView,
    ReportKey,
    ReportOptionsView,
    ReportPartnerOption,
    ReportView,
)
from app.modules.returns.models import ReturnRefund
from app.modules.returns.service import invoice_net_amounts
from app.modules.sales.models import (
    CustomerInvoice,
    SalesOrder,
    SalesOrderLine,
    SalesQuotation,
    SalesQuotationLine,
    SalesWeightCard,
    SalesWeightCardLine,
)
from app.modules.treasury.models import (
    FinancialAccount,
    PaymentAllocation,
    PaymentTransaction,
)
from app.modules.treasury.schemas import StatementLineView
from app.modules.treasury.service import list_partner_balances, partner_statement

ZERO = Decimal("0")
REPORT_TITLES: dict[str, str] = {
    "sales": "تقرير المبيعات",
    "purchases": "تقرير المشتريات",
    "customer_balances": "أرصدة العملاء",
    "supplier_balances": "أرصدة الموردين",
    "payments": "حركات التحصيل والسداد",
    "inventory_valuation": "تقييم المخزون",
}


class ReportError(Exception):
    pass


class PrintDocumentNotFound(ReportError):
    pass


def _money(value: Decimal | int | str) -> str:
    return f"{money(Decimal(str(value))):.2f}"


def _quantity(value: Decimal | int | str) -> str:
    return f"{quantity(Decimal(str(value))):.6f}"


def _window(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    if date_from > date_to:
        raise ReportError("تاريخ البداية يجب ألا يكون بعد تاريخ النهاية")
    return (
        datetime.combine(date_from, time.min, tzinfo=UTC),
        datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=UTC),
    )


def report_options(db: Session) -> ReportOptionsView:
    partners = db.scalars(
        select(Partner).where(Partner.is_active.is_(True)).order_by(Partner.name_ar)
    )
    return ReportOptionsView(
        partners=[
            ReportPartnerOption(
                id=item.id,
                code=item.code,
                name_ar=item.name_ar,
                is_customer=item.is_customer,
                is_supplier=item.is_supplier,
            )
            for item in partners
        ]
    )


def _company(db: Session) -> PrintCompanyView:
    settings = db.get(CompanySettings, 1)
    if settings is None:
        raise PrintDocumentNotFound("إعدادات الشركة غير موجودة")
    return PrintCompanyView(
        name_ar=settings.company_name_ar,
        phone=settings.phone,
        address=settings.address,
        tax_number=settings.tax_number,
        currency_code=settings.currency_code,
    )


def _payment_methods(db: Session, invoice_id: UUID) -> list[str]:
    direct = db.scalars(
        select(PaymentTransaction.payment_method).where(
            PaymentTransaction.customer_invoice_id == invoice_id,
            PaymentTransaction.status == "posted",
        )
    )
    allocated = db.scalars(
        select(PaymentTransaction.payment_method)
        .join(
            PaymentAllocation,
            PaymentAllocation.payment_transaction_id == PaymentTransaction.id,
        )
        .where(
            PaymentAllocation.customer_invoice_id == invoice_id,
            PaymentTransaction.status == "posted",
        )
    )
    return sorted(set(direct) | set(allocated))


def sales_invoice_print_data(db: Session, invoice_id: UUID) -> PrintDocumentView:
    row = db.execute(
        select(CustomerInvoice, SalesOrder, Partner)
        .join(SalesOrder, SalesOrder.id == CustomerInvoice.sales_order_id)
        .join(Partner, Partner.id == CustomerInvoice.customer_id)
        .where(CustomerInvoice.id == invoice_id)
    ).one_or_none()
    if row is None:
        raise PrintDocumentNotFound("فاتورة المبيعات غير موجودة")
    invoice, order, partner = row
    if invoice.status != "posted":
        raise ReportError("يمكن طباعة فاتورة مبيعات معتمدة فقط")
    returned, paid, refunded, remaining, _ = invoice_net_amounts(
        db,
        return_type="sales",
        invoice_id=invoice.id,
        original_total=invoice.total,
    )
    net = money(max(ZERO, invoice.total - returned))
    order_lines = list(
        db.scalars(
            select(SalesOrderLine)
            .where(SalesOrderLine.sales_order_id == order.id)
            .order_by(SalesOrderLine.created_at, SalesOrderLine.id)
        )
    )
    products = {
        item.id: item
        for item in db.scalars(
            select(Product).where(Product.id.in_({line.product_id for line in order_lines}))
        )
    }
    card = db.scalar(select(SalesWeightCard).where(SalesWeightCard.sales_order_id == order.id))
    weight_lines = (
        {
            item.sales_order_line_id: item
            for item in db.scalars(
                select(SalesWeightCardLine).where(SalesWeightCardLine.weight_card_id == card.id)
            )
        }
        if card is not None
        else {}
    )
    document_type = "weight_invoice" if invoice.invoice_type == "weight" else "sales_invoice"
    return PrintDocumentView(
        document_type=cast(Literal["sales_invoice", "weight_invoice"], document_type),
        document_title=(
            "فاتورة مبيعات بالوزن / الكارتة"
            if document_type == "weight_invoice"
            else "فاتورة مبيعات"
        ),
        document_number=invoice.invoice_number,
        document_date=invoice.invoice_date,
        order_number=order.order_number,
        partner_name_ar=partner.name_ar,
        partner_code=partner.code,
        partner_phone=partner.phone,
        partner_address=partner.address,
        status=invoice.status,
        payment_methods=_payment_methods(db, invoice.id),
        subtotal=invoice.subtotal,
        discount_amount=invoice.discount_amount,
        transport_amount=invoice.transport_amount,
        tax_amount=invoice.tax_amount,
        original_total=invoice.total,
        returned_total=returned,
        net_total=net,
        paid=paid,
        refunded=refunded,
        remaining=remaining,
        notes=invoice.notes,
        card_number=card.card_number if card is not None else "",
        vehicle_number=card.vehicle_number if card is not None else "",
        gross_weight_kg=card.gross_weight_kg if card is not None else ZERO,
        tare_weight_kg=card.tare_weight_kg if card is not None else ZERO,
        net_weight_kg=card.net_weight_kg if card is not None else ZERO,
        lines=[
            PrintDocumentLineView(
                code=products[line.product_id].code,
                name=products[line.product_id].name_ar,
                quantity=line.quantity,
                unit=line.unit,
                unit_price=line.unit_price,
                line_total=line.line_total,
                notes=line.notes,
                actual_weight_kg=(
                    weight_lines[line.id].actual_weight_kg if line.id in weight_lines else None
                ),
                price_per_kg=(
                    weight_lines[line.id].price_per_kg if line.id in weight_lines else None
                ),
            )
            for line in order_lines
        ],
        company=_company(db),
    )


def quotation_print_data(db: Session, quotation_id: UUID) -> PrintDocumentView:
    row = db.execute(
        select(SalesQuotation, Partner)
        .join(Partner, Partner.id == SalesQuotation.customer_id)
        .where(SalesQuotation.id == quotation_id)
    ).one_or_none()
    if row is None:
        raise PrintDocumentNotFound("عرض السعر غير موجود")
    quotation, partner = row
    lines = list(
        db.scalars(
            select(SalesQuotationLine)
            .where(SalesQuotationLine.quotation_id == quotation.id)
            .order_by(SalesQuotationLine.created_at, SalesQuotationLine.id)
        )
    )
    products = {
        item.id: item
        for item in db.scalars(
            select(Product).where(
                Product.id.in_({line.product_id for line in lines if line.product_id is not None})
            )
        )
    }
    return PrintDocumentView(
        document_type="quotation",
        document_title="عرض سعر",
        document_number=quotation.quotation_number,
        document_date=quotation.quotation_date,
        order_number="—",
        partner_name_ar=partner.name_ar,
        partner_code=partner.code,
        partner_phone=partner.phone,
        partner_address=partner.address,
        status=quotation.status,
        payment_methods=[],
        subtotal=quotation.total,
        discount_amount=ZERO,
        transport_amount=ZERO,
        tax_amount=ZERO,
        original_total=quotation.total,
        returned_total=ZERO,
        net_total=quotation.total,
        paid=ZERO,
        refunded=ZERO,
        remaining=quotation.total,
        notes=quotation.notes,
        valid_until=quotation.valid_until,
        lines=[
            PrintDocumentLineView(
                code=(products[line.product_id].code if line.product_id in products else ""),
                name=line.item_name,
                quantity=line.quantity,
                unit=line.unit,
                unit_price=line.unit_price,
                line_total=line.line_total,
                notes=line.notes,
            )
            for line in lines
        ],
        company=_company(db),
    )


def customer_statement_print_data(
    db: Session,
    *,
    customer_id: UUID,
    date_from: date,
    date_to: date,
    detailed: bool = False,
    include_drafts: bool = False,
) -> CustomerStatementPrintView:
    customer = db.get(Partner, customer_id)
    if customer is None or not customer.is_customer:
        raise PrintDocumentNotFound("العميل غير موجود")
    statement = partner_statement(
        db,
        partner_id=customer_id,
        date_from=date_from,
        date_to=date_to,
        partner_type="customer",
    )
    start, end = _window(date_from, date_to)
    if include_drafts:
        draft_rows = db.scalars(
            select(SalesOrder)
            .where(
                SalesOrder.customer_id == customer_id,
                SalesOrder.status == "draft",
                SalesOrder.order_date >= start,
                SalesOrder.order_date < end,
            )
            .order_by(SalesOrder.order_date, SalesOrder.id)
        )
        merged = list(statement.lines)
        merged.extend(
            StatementLineView(
                movement_date=order.order_date,
                document_number=order.order_number,
                movement_type=(
                    "مسودة بيع بالوزن" if order.billing_method == "weight" else "مسودة بيع عادية"
                ),
                debit=ZERO,
                credit=ZERO,
                running_balance=ZERO,
                notes=f"للمراجعة فقط — لا تدخل في الرصيد — قيمة {_money(order.total)}",
            )
            for order in draft_rows
        )
        merged.sort(
            key=lambda item: (
                item.movement_date
                if item.movement_date.tzinfo is not None
                else item.movement_date.replace(tzinfo=UTC),
                item.document_number,
            )
        )
        running = statement.opening_balance
        recalculated: list[StatementLineView] = []
        for line in merged:
            running = money(running + line.debit - line.credit)
            recalculated.append(line.model_copy(update={"running_balance": running}))
        statement = statement.model_copy(update={"lines": recalculated, "closing_balance": running})
    invoice_details: dict[str, list[PrintDocumentLineView]] = {}
    if detailed:
        invoice_rows = db.execute(
            select(CustomerInvoice, SalesOrderLine, Product)
            .join(SalesOrder, SalesOrder.id == CustomerInvoice.sales_order_id)
            .join(SalesOrderLine, SalesOrderLine.sales_order_id == SalesOrder.id)
            .join(Product, Product.id == SalesOrderLine.product_id)
            .where(
                CustomerInvoice.customer_id == customer_id,
                CustomerInvoice.status == "posted",
                CustomerInvoice.invoice_date >= start,
                CustomerInvoice.invoice_date < end,
            )
            .order_by(CustomerInvoice.invoice_date, SalesOrderLine.created_at)
        )
        for invoice, line, product in invoice_rows:
            invoice_details.setdefault(invoice.invoice_number, []).append(
                PrintDocumentLineView(
                    code=product.code,
                    name=product.name_ar,
                    quantity=line.quantity,
                    unit=line.unit,
                    unit_price=line.unit_price,
                    line_total=line.line_total,
                    notes=line.notes,
                )
            )
    invoice_totals: dict[str, Decimal] = {}
    invoice_total_rows = db.execute(
        select(
            CustomerInvoice.invoice_type,
            func.coalesce(func.sum(CustomerInvoice.total), 0),
        )
        .where(
            CustomerInvoice.customer_id == customer_id,
            CustomerInvoice.status == "posted",
            CustomerInvoice.invoice_date >= start,
            CustomerInvoice.invoice_date < end,
        )
        .group_by(CustomerInvoice.invoice_type)
    )
    for invoice_type, total in invoice_total_rows.tuples():
        invoice_totals[invoice_type] = money(total)
    returns_total = ZERO
    receipts_total = ZERO
    refunds_total = ZERO
    adjustments_total = ZERO
    for line in statement.lines:
        if line.movement_type == "مرتجع مبيعات":
            returns_total += line.credit
        elif line.movement_type == "تحصيل عميل":
            receipts_total += line.credit
        elif line.movement_type == "رد مبلغ لعميل":
            refunds_total += line.debit
        elif line.movement_type == "تسوية حساب عميل":
            adjustments_total += line.debit - line.credit
    summary = CustomerStatementSummaryView(
        opening_balance=statement.opening_balance,
        standard_sales_total=money(invoice_totals.get("standard", ZERO)),
        weight_sales_total=money(invoice_totals.get("weight", ZERO)),
        returns_total=money(returns_total),
        receipts_total=money(receipts_total),
        customer_refunds_total=money(refunds_total),
        adjustments_total=money(adjustments_total),
        net_movement=money(
            sum((line.debit - line.credit for line in statement.lines), ZERO)
        ),
        closing_balance=statement.closing_balance,
    )
    return CustomerStatementPrintView(
        company=_company(db),
        partner_phone=customer.phone,
        statement=statement,
        detailed=detailed,
        include_drafts=include_drafts,
        invoice_details=invoice_details,
        summary=summary,
    )


def customer_statement_xlsx(view: CustomerStatementPrintView) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "كشف الحساب"
    sheet.sheet_view.rightToLeft = True
    statement = view.statement
    sheet.append(["كشف حساب العميل", statement.partner_name_ar])
    sheet.append(["من", statement.date_from.isoformat(), "إلى", statement.date_to.isoformat()])
    sheet.append(
        ["رصيد أول المدة", statement.opening_balance, "الرصيد النهائي", statement.closing_balance]
    )
    sheet.append([])
    headers = ["التاريخ", "المستند", "النوع / البيان", "مدين", "دائن", "الرصيد"]
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for line in statement.lines:
        sheet.append(
            [
                line.movement_date.isoformat(),
                line.document_number,
                f"{line.movement_type} {line.notes}".strip(),
                line.debit,
                line.credit,
                line.running_balance,
            ]
        )
        for detail in view.invoice_details.get(line.document_number, []):
            sheet.append(
                [
                    "",
                    "",
                    f"↳ {detail.code} · {detail.name} · {detail.quantity} {detail.unit}",
                    "",
                    "",
                    detail.line_total,
                ]
            )
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:F{max(5, sheet.max_row)}"
    for cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in cells)
        sheet.column_dimensions[cells[0].column_letter].width = min(width + 4, 55)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _sales_report(
    db: Session, *, start: datetime, end: datetime, partner_id: UUID | None
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    statement = (
        select(SalesOrder, CustomerInvoice, Partner)
        .join(Partner, Partner.id == SalesOrder.customer_id)
        .outerjoin(CustomerInvoice, CustomerInvoice.sales_order_id == SalesOrder.id)
        .where(SalesOrder.order_date >= start, SalesOrder.order_date < end)
        .order_by(SalesOrder.order_date.desc(), SalesOrder.id.desc())
    )
    if partner_id is not None:
        statement = statement.where(SalesOrder.customer_id == partner_id)
    rows: list[dict[str, str]] = []
    totals = {"original": ZERO, "returned": ZERO, "net": ZERO, "paid": ZERO, "remaining": ZERO}
    for order, invoice, partner in db.execute(statement):
        original = money(invoice.total if invoice is not None else order.total)
        returned = paid = refunded = remaining = ZERO
        if invoice is not None and invoice.status == "posted":
            returned, paid, refunded, remaining, _ = invoice_net_amounts(
                db,
                return_type="sales",
                invoice_id=invoice.id,
                original_total=invoice.total,
            )
        net = money(max(ZERO, original - returned))
        effective_paid = money(max(ZERO, paid - refunded))
        rows.append(
            {
                "رقم الأمر": order.order_number,
                "التاريخ": order.order_date.date().isoformat(),
                "العميل": partner.name_ar,
                "الإجمالي الأصلي": _money(original),
                "المرتجع": _money(returned),
                "الصافي": _money(net),
                "المدفوع": _money(effective_paid),
                "المتبقي": _money(remaining if invoice is not None else net),
                "الحالة": order.status,
            }
        )
        totals["original"] += original
        totals["returned"] += returned
        totals["net"] += net
        totals["paid"] += effective_paid
        totals["remaining"] += remaining if invoice is not None else net
    columns = [
        "رقم الأمر",
        "التاريخ",
        "العميل",
        "الإجمالي الأصلي",
        "المرتجع",
        "الصافي",
        "المدفوع",
        "المتبقي",
        "الحالة",
    ]
    return columns, rows, {key: _money(value) for key, value in totals.items()}


def _purchases_report(
    db: Session, *, start: datetime, end: datetime, partner_id: UUID | None
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    statement = (
        select(PurchaseOrder, SupplierInvoice, Partner)
        .join(Partner, Partner.id == PurchaseOrder.supplier_id)
        .outerjoin(SupplierInvoice, SupplierInvoice.purchase_order_id == PurchaseOrder.id)
        .where(PurchaseOrder.order_date >= start, PurchaseOrder.order_date < end)
        .order_by(PurchaseOrder.order_date.desc(), PurchaseOrder.id.desc())
    )
    if partner_id is not None:
        statement = statement.where(PurchaseOrder.supplier_id == partner_id)
    rows: list[dict[str, str]] = []
    totals = {"original": ZERO, "returned": ZERO, "net": ZERO, "paid": ZERO, "remaining": ZERO}
    for order, invoice, partner in db.execute(statement):
        original = money(invoice.total if invoice is not None else order.total)
        returned = paid = refunded = remaining = ZERO
        if invoice is not None and invoice.status == "posted":
            returned, paid, refunded, remaining, _ = invoice_net_amounts(
                db,
                return_type="purchase",
                invoice_id=invoice.id,
                original_total=invoice.total,
            )
        net = money(max(ZERO, original - returned))
        effective_paid = money(max(ZERO, paid - refunded))
        rows.append(
            {
                "رقم الأمر": order.order_number,
                "التاريخ": order.order_date.date().isoformat(),
                "المورد": partner.name_ar,
                "الإجمالي الأصلي": _money(original),
                "المرتجع": _money(returned),
                "الصافي": _money(net),
                "المدفوع": _money(effective_paid),
                "المتبقي": _money(remaining if invoice is not None else net),
                "الحالة": order.status,
            }
        )
        totals["original"] += original
        totals["returned"] += returned
        totals["net"] += net
        totals["paid"] += effective_paid
        totals["remaining"] += remaining if invoice is not None else net
    columns = [
        "رقم الأمر",
        "التاريخ",
        "المورد",
        "الإجمالي الأصلي",
        "المرتجع",
        "الصافي",
        "المدفوع",
        "المتبقي",
        "الحالة",
    ]
    return columns, rows, {key: _money(value) for key, value in totals.items()}


def _balances_report(
    db: Session, *, partner_type: str, partner_id: UUID | None
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    name_label = "العميل" if partner_type == "customer" else "المورد"
    values = list_partner_balances(db, partner_type=partner_type)
    if partner_id is not None:
        values = [item for item in values if item.partner_id == partner_id]
    rows = [
        {
            name_label: item.partner_name_ar,
            "الرصيد الافتتاحي": _money(item.opening_balance),
            "صافي الفواتير": _money(item.invoices_total),
            "المرتجعات": _money(item.returns_total),
            "المدفوع": _money(item.paid_total),
            "الاستردادات": _money(item.refunds_total),
            "دفعات مقدمة": _money(item.advances),
            "التسويات": _money(item.adjustments_total),
            "الرصيد": _money(item.balance),
        }
        for item in values
    ]
    total_balance = sum((item.balance for item in values), ZERO)
    columns = [
        name_label,
        "الرصيد الافتتاحي",
        "صافي الفواتير",
        "المرتجعات",
        "المدفوع",
        "الاستردادات",
        "دفعات مقدمة",
        "التسويات",
        "الرصيد",
    ]
    return columns, rows, {"balance": _money(total_balance)}


def _payments_report(
    db: Session, *, start: datetime, end: datetime, partner_id: UUID | None
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    payment_statement = (
        select(PaymentTransaction, Partner, FinancialAccount)
        .join(Partner, Partner.id == PaymentTransaction.partner_id)
        .join(FinancialAccount, FinancialAccount.id == PaymentTransaction.financial_account_id)
        .where(
            PaymentTransaction.transaction_date >= start, PaymentTransaction.transaction_date < end
        )
    )
    refund_statement = (
        select(ReturnRefund, Partner, FinancialAccount)
        .join(Partner, Partner.id == ReturnRefund.partner_id)
        .join(FinancialAccount, FinancialAccount.id == ReturnRefund.financial_account_id)
        .where(ReturnRefund.refund_date >= start, ReturnRefund.refund_date < end)
    )
    if partner_id is not None:
        payment_statement = payment_statement.where(PaymentTransaction.partner_id == partner_id)
        refund_statement = refund_statement.where(ReturnRefund.partner_id == partner_id)
    values: list[tuple[datetime, dict[str, str], Decimal]] = []
    for item, partner, account in db.execute(payment_statement):
        label = "تحصيل عميل" if item.transaction_type == "customer_receipt" else "سداد مورد"
        values.append(
            (
                item.transaction_date,
                {
                    "رقم الحركة": item.transaction_number,
                    "التاريخ": item.transaction_date.isoformat(),
                    "النوع": label,
                    "الطرف": partner.name_ar,
                    "الحساب": account.name_ar,
                    "المبلغ": _money(item.amount),
                    "الطريقة": item.payment_method,
                    "الحالة": item.status,
                    "ملاحظات": item.notes,
                },
                money(item.amount),
            )
        )
    for item, partner, account in db.execute(refund_statement):
        label = "رد مبلغ لعميل" if item.refund_type == "customer_refund" else "استرداد من مورد"
        values.append(
            (
                item.refund_date,
                {
                    "رقم الحركة": item.refund_number,
                    "التاريخ": item.refund_date.isoformat(),
                    "النوع": label,
                    "الطرف": partner.name_ar,
                    "الحساب": account.name_ar,
                    "المبلغ": _money(item.amount),
                    "الطريقة": item.payment_method,
                    "الحالة": item.status,
                    "ملاحظات": item.notes,
                },
                money(item.amount),
            )
        )
    values.sort(key=lambda value: value[0], reverse=True)
    columns = [
        "رقم الحركة",
        "التاريخ",
        "النوع",
        "الطرف",
        "الحساب",
        "المبلغ",
        "الطريقة",
        "الحالة",
        "ملاحظات",
    ]
    return (
        columns,
        [item[1] for item in values],
        {"total": _money(sum((item[2] for item in values), ZERO))},
    )


def _inventory_report(
    db: Session,
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    balances = db.execute(
        select(InventoryBalance, Product, Warehouse)
        .join(Product, Product.id == InventoryBalance.product_id)
        .join(Warehouse, Warehouse.id == InventoryBalance.warehouse_id)
        .where(Product.is_active.is_(True))
        .order_by(Product.name_ar, Warehouse.name_ar)
    )
    rows: list[dict[str, str]] = []
    total_value = ZERO
    for balance, product, warehouse in balances:
        value = db.scalar(
            select(
                func.coalesce(
                    func.sum(
                        case(
                            (
                                InventoryLayer.cost_basis == "weight",
                                InventoryLayer.weight_remaining_kg * InventoryLayer.unit_cost,
                            ),
                            else_=InventoryLayer.quantity_remaining * InventoryLayer.unit_cost,
                        )
                    ),
                    0,
                )
            ).where(
                InventoryLayer.product_id == product.id, InventoryLayer.warehouse_id == warehouse.id
            )
        )
        layer_value = money(Decimal(str(value or 0)))
        basis_amount = (
            balance.weight_on_hand_kg if balance.weight_on_hand_kg > 0 else balance.quantity_on_hand
        )
        average = money(layer_value / basis_amount) if basis_amount > 0 else ZERO
        total_value += layer_value
        rows.append(
            {
                "الكود": product.code,
                "الصنف": product.name_ar,
                "المخزن": warehouse.name_ar,
                "الكمية": _quantity(balance.quantity_on_hand),
                "الوزن كجم": _quantity(balance.weight_on_hand_kg),
                "متوسط التكلفة": _money(average),
                "قيمة المخزون": _money(layer_value),
            }
        )
    columns = ["الكود", "الصنف", "المخزن", "الكمية", "الوزن كجم", "متوسط التكلفة", "قيمة المخزون"]
    return columns, rows, {"inventory_value": _money(total_value)}


def generate_report(
    db: Session,
    *,
    report_key: str,
    date_from: date,
    date_to: date,
    partner_id: UUID | None = None,
) -> ReportView:
    start, end = _window(date_from, date_to)
    if report_key == "sales":
        columns, rows, summary = _sales_report(db, start=start, end=end, partner_id=partner_id)
    elif report_key == "purchases":
        columns, rows, summary = _purchases_report(db, start=start, end=end, partner_id=partner_id)
    elif report_key in {"customer_balances", "supplier_balances"}:
        columns, rows, summary = _balances_report(
            db,
            partner_type="customer" if report_key == "customer_balances" else "supplier",
            partner_id=partner_id,
        )
    elif report_key == "payments":
        columns, rows, summary = _payments_report(db, start=start, end=end, partner_id=partner_id)
    elif report_key == "inventory_valuation":
        columns, rows, summary = _inventory_report(db)
    else:
        raise ReportError("نوع التقرير غير مدعوم")
    return ReportView(
        report_key=cast(ReportKey, report_key),
        title=REPORT_TITLES[report_key],
        date_from=date_from,
        date_to=date_to,
        columns=columns,
        rows=rows,
        summary=summary,
    )


def report_xlsx(view: ReportView) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "التقرير"
    sheet.sheet_view.rightToLeft = True
    sheet.append([view.title])
    sheet.append(["من تاريخ", view.date_from.isoformat(), "إلى تاريخ", view.date_to.isoformat()])
    sheet.append([])
    sheet.append(view.columns)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in view.rows:
        sheet.append([row.get(column, "") for column in view.columns])
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = (
        f"A{header_row}:{sheet.cell(max(header_row, sheet.max_row), len(view.columns)).coordinate}"
    )
    for cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in cells)
        sheet.column_dimensions[cells[0].column_letter].width = min(width + 4, 45)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
